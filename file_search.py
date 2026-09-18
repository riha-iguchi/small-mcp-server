"""ローカルフォルダ内のファイルを検索・取得するためのヘルパーモジュール。

Teamsからエクスポートしたマニュアルなどを`DOCS_DIR`配下に置いておくと、
その中身をキーワード検索したり、全文を取得したりできる。
"""
from __future__ import annotations

import io
import os
from pathlib import Path

DOCS_DIR = Path(os.environ.get("DOCS_DIR", "docs")).resolve()

# 検索・読み取りの対象とする拡張子
SUPPORTED_EXTENSIONS = {".docx", ".xlsx", ".pdf", ".txt", ".md"}


def _resolve_path(relative_path: str) -> Path:
    """DOCS_DIR配下のパスに解決する。範囲外(親ディレクトリへの脱出)は拒否する。"""
    candidate = (DOCS_DIR / relative_path).resolve()
    if DOCS_DIR not in candidate.parents and candidate != DOCS_DIR:
        raise ValueError(f"'{relative_path}' は許可されたフォルダの外を指しています。")
    return candidate


def _extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    content = path.read_bytes()

    if ext in (".txt", ".md"):
        for enc in ("utf-8-sig", "utf-8", "cp932"):
            try:
                return content.decode(enc)
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="replace")

    if ext == ".docx":
        import docx

        doc = docx.Document(io.BytesIO(content))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append("\t".join(cell.text for cell in row.cells))
        return "\n".join(parts)

    if ext == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        parts = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append("\t".join(cells))
        return "\n".join(parts)

    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    raise ValueError(f"未対応のファイル形式です: {ext}")


def list_files(folder: str = "") -> list[dict[str, str]]:
    """DOCS_DIR配下(folder指定時はそのサブフォルダ)のファイル・フォルダ一覧を返す。"""
    target = _resolve_path(folder)
    if not target.exists():
        raise FileNotFoundError(f"フォルダが見つかりません: {folder or '(ルート)'}")
    items = []
    for entry in sorted(target.iterdir()):
        items.append({"name": entry.name, "type": "folder" if entry.is_dir() else "file"})
    return items


def read_file(file_path: str) -> str:
    """指定したファイル1件のテキスト内容を取得する。"""
    path = _resolve_path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"ファイルが見つかりません: {file_path}")
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"未対応のファイル形式です: {path.suffix}")
    return _extract_text(path)


def search_files(query: str, folder: str = "", max_results: int = 5) -> list[dict[str, str]]:
    """folder配下のファイルを再帰的に検索し、queryに一致した箇所を返す(単純な部分一致)。"""
    target = _resolve_path(folder)
    if not target.exists():
        raise FileNotFoundError(f"フォルダが見つかりません: {folder or '(ルート)'}")

    query_lower = query.lower()
    results: list[dict[str, str]] = []
    for path in sorted(target.rglob("*")):
        if len(results) >= max_results:
            break
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        try:
            text = _extract_text(path)
        except Exception:
            continue
        idx = text.lower().find(query_lower)
        if idx == -1:
            continue
        start = max(0, idx - 150)
        end = min(len(text), idx + len(query) + 150)
        rel = path.relative_to(DOCS_DIR).as_posix()
        results.append({"file": rel, "snippet": text[start:end].strip()})
    return results
